from pathlib import Path

import pytest
from docx import Document
from pypdf import PdfWriter

from medical_journal_matcher.parsing import parse_manuscript


@pytest.mark.parametrize("suffix", [".txt", ".md"])
def test_parse_plain_text_formats(tmp_path: Path, suffix: str) -> None:
    manuscript = tmp_path / f"manuscript{suffix}"
    manuscript.write_text("title\nabstract", encoding="utf-8")
    assert parse_manuscript(manuscript) == "title\nabstract"


def test_parse_docx(tmp_path: Path) -> None:
    manuscript = tmp_path / "manuscript.docx"
    document = Document()
    document.add_paragraph("Title")
    document.add_paragraph("Abstract")
    document.save(manuscript)
    assert parse_manuscript(manuscript) == "Title\nAbstract"


def test_parse_blank_pdf(tmp_path: Path) -> None:
    manuscript = tmp_path / "manuscript.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    with manuscript.open("wb") as stream:
        writer.write(stream)
    assert parse_manuscript(manuscript) == ""


def test_parse_rejects_unsupported_format(tmp_path: Path) -> None:
    manuscript = tmp_path / "manuscript.rtf"
    manuscript.write_text("text", encoding="utf-8")
    with pytest.raises(ValueError, match="Supported manuscript formats"):
        parse_manuscript(manuscript)
